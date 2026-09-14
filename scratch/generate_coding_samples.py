import os
import json
import csv
import zipfile
import docx

SAMPLES_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'static', 'samples')
os.makedirs(SAMPLES_DIR, exist_ok=True)

DEFAULT_BOILERPLATES = {
    'python': "a, b = map(int, input().split())\nprint(a + b)",
    'cpp': "#include <iostream>\nusing namespace std;\nint main() {\n    int a, b;\n    if (cin >> a >> b) cout << a + b << endl;\n    return 0;\n}",
    'c': "#include <stdio.h>\nint main() {\n    int a, b;\n    if (scanf(\"%d %d\", &a, &b) == 2) printf(\"%d\\n\", a + b);\n    return 0;\n}",
    'java': "import java.util.*;\npublic class Main {\n    public static void main(String[] args) {\n        Scanner sc = new Scanner(System.in);\n        if (sc.hasNextInt()) {\n            int a = sc.nextInt(), b = sc.nextInt();\n            System.out.println(a + b);\n        }\n    }\n}",
    'javascript': "const fs = require('fs');\nconst [a, b] = fs.readFileSync(0, 'utf8').trim().split(/\\s+/).map(Number);\nconsole.log(a + b);"
}

# 1. JSON
json_path = os.path.join(SAMPLES_DIR, 'sample_coding.json')
sample_json_data = {
    "assessmentId": "assessment_001",
    "questions": [
        {
            "id": "code_001",
            "type": "coding",
            "title": "Add Two Numbers",
            "difficulty": "Easy",
            "category": "Basic Programming",
            "tags": ["arithmetic", "input-output"],
            "problemStatement": "Write a program that reads two integers and prints their sum.",
            "inputFormat": "Two space-separated integers.",
            "outputFormat": "Print the sum of the two integers.",
            "constraints": ["-1000000000 <= a, b <= 1000000000"],
            "supportedLanguages": ["python", "c", "cpp", "java", "javascript"],
            "starterCode": DEFAULT_BOILERPLATES,
            "samples": [
                {
                    "input": "5 7",
                    "expectedOutput": "12",
                    "explanation": "5 + 7 = 12"
                }
            ],
            "testCases": [
                {
                    "id": "tc_001",
                    "input": "5 7",
                    "expectedOutput": "12",
                    "hidden": False,
                    "weight": 20
                },
                {
                    "id": "tc_002",
                    "input": "10 20",
                    "expectedOutput": "30",
                    "hidden": True,
                    "weight": 40
                },
                {
                    "id": "tc_003",
                    "input": "-5 8",
                    "expectedOutput": "3",
                    "hidden": True,
                    "weight": 40
                }
            ],
            "marks": 10,
            "negativeMarks": 2,
            "timeLimitMs": 2000,
            "memoryLimitMb": 256,
            "status": "draft",
            "explanation": "The program should read two integers and print their sum."
        },
        {
            "id": "code_002",
            "type": "coding",
            "title": "Reverse a String",
            "difficulty": "Medium",
            "category": "Strings",
            "tags": ["string", "two-pointers"],
            "problemStatement": "Given a string S, output the reverse of the string.",
            "inputFormat": "A single line containing the string S.",
            "outputFormat": "Print the reversed string.",
            "constraints": ["1 <= length(S) <= 10^5"],
            "supportedLanguages": ["python", "c", "cpp", "java", "javascript"],
            "starterCode": DEFAULT_BOILERPLATES,
            "samples": [
                {
                    "input": "hello",
                    "expectedOutput": "olleh",
                    "explanation": "Reversing 'hello' gives 'olleh'"
                }
            ],
            "testCases": [
                {
                    "id": "tc_001",
                    "input": "hello",
                    "expectedOutput": "olleh",
                    "hidden": False,
                    "weight": 50
                },
                {
                    "id": "tc_002",
                    "input": "racecar",
                    "expectedOutput": "racecar",
                    "hidden": True,
                    "weight": 50
                }
            ],
            "marks": 15,
            "negativeMarks": 3,
            "timeLimitMs": 2000,
            "memoryLimitMb": 256,
            "status": "published",
            "explanation": "Reverse the string."
        }
    ]
}
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(sample_json_data, f, indent=2)
print("Created:", json_path)

# 2. CSV
csv_path = os.path.join(SAMPLES_DIR, 'sample_coding.csv')
fieldnames = [
    'id', 'type', 'title', 'difficulty', 'category', 'tags',
    'problemStatement', 'inputFormat', 'outputFormat', 'constraints',
    'supportedLanguages', 'starterCode', 'samples', 'testCases',
    'marks', 'negativeMarks', 'timeLimitMs', 'memoryLimitMb', 'status'
]
rows = [
    {
        'id': 'code_001',
        'type': 'coding',
        'title': 'Add Two Numbers',
        'difficulty': 'Easy',
        'category': 'Basic Programming',
        'tags': 'arithmetic,input-output',
        'problemStatement': 'Write a program that reads two integers and prints their sum.',
        'inputFormat': 'Two space-separated integers.',
        'outputFormat': 'Print the sum of the two integers.',
        'constraints': json.dumps(["-1000000000 <= a, b <= 1000000000"]),
        'supportedLanguages': 'python,c,cpp,java,javascript',
        'starterCode': json.dumps(DEFAULT_BOILERPLATES),
        'samples': json.dumps([{"input": "5 7", "expectedOutput": "12", "explanation": "5 + 7 = 12"}]),
        'testCases': json.dumps([
            {"id": "tc_001", "input": "5 7", "expectedOutput": "12", "hidden": False, "weight": 20},
            {"id": "tc_002", "input": "10 20", "expectedOutput": "30", "hidden": True, "weight": 40},
            {"id": "tc_003", "input": "-5 8", "expectedOutput": "3", "hidden": True, "weight": 40}
        ]),
        'marks': 10,
        'negativeMarks': 2,
        'timeLimitMs': 2000,
        'memoryLimitMb': 256,
        'status': 'draft'
    }
]
with open(csv_path, 'w', encoding='utf-8', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
print("Created:", csv_path)

# 3. DOCX
docx_path = os.path.join(SAMPLES_DIR, 'sample_coding.docx')
doc = docx.Document()
doc.add_heading('QuizMaster Coding Question Template', level=1)
doc.add_paragraph('Fill in the table below to import coding questions. Each table represents one coding challenge.')

table = doc.add_table(rows=15, cols=2)
table.style = 'Table Grid'
fields = [
    ('Question ID', 'code_001'),
    ('Title', 'Add Two Numbers'),
    ('Difficulty', 'Easy'),
    ('Category', 'Basic Programming'),
    ('Problem Statement', 'Write a program that reads two integers and prints their sum.'),
    ('Input Format', 'Two space-separated integers.'),
    ('Output Format', 'Print the sum of the two integers.'),
    ('Constraints', '-10^9 <= a, b <= 10^9'),
    ('Supported Languages', 'python, c, cpp, java, javascript'),
    ('Starter Code', json.dumps(DEFAULT_BOILERPLATES)),
    ('Sample Input', '5 7'),
    ('Sample Output', '12'),
    ('Test Cases', json.dumps([
        {"id": "tc_001", "input": "5 7", "expectedOutput": "12", "hidden": False, "weight": 50},
        {"id": "tc_002", "input": "10 20", "expectedOutput": "30", "hidden": True, "weight": 50}
    ])),
    ('Marks', '10'),
    ('Negative Marks', '2')
]
for i, (k, v) in enumerate(fields):
    table.rows[i].cells[0].text = k
    table.rows[i].cells[1].text = v

doc.save(docx_path)
print("Created:", docx_path)

# 4. ZIP
zip_path = os.path.join(SAMPLES_DIR, 'sample_coding.zip')
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
    zf.writestr('questions.json', json.dumps(sample_json_data, indent=2))
    zf.writestr('code_001_tc1.in', '5 7')
    zf.writestr('code_001_tc1.out', '12')
    zf.writestr('code_001_tc2.in', '10 20')
    zf.writestr('code_001_tc2.out', '30')
print("Created:", zip_path)
print("All sample templates created successfully!")
