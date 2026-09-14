from flask import Flask, request, redirect, session, render_template, jsonify, flash, send_file, Response, current_app, url_for
from datetime import timedelta, datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import random
import csv
import io
import base64
import time
import threading
import json
import re
from dotenv import load_dotenv
from database import get_db_connection
from certificate_generator import generate_certificate_pdf
from utils import *

load_dotenv()
app = Flask(__name__)

# Production Secret Key Handling
flask_secret = os.getenv('FLASK_SECRET_KEY')
if not flask_secret or flask_secret in ('change-me', 'dev-secret-change-me'):
    flask_secret = os.getenv('FLASK_SECRET_KEY', 'qm-prod-' + secrets.token_hex(24))
app.secret_key = flask_secret

is_dev = os.getenv('FLASK_DEBUG', '0') == '1'
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=1)
app.config['SESSION_COOKIE_SECURE'] = os.getenv('SESSION_COOKIE_SECURE', '0') == '1'
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['TEMPLATES_AUTO_RELOAD'] = is_dev
app.jinja_env.auto_reload = is_dev

# Blueprint Registration
from routes.auth import auth_bp
from routes.home import home_bp
from routes.admin import admin_bp
from routes.coordinator import coordinator_bp
from routes.quiz import quiz_bp
from routes.study_materials import study_bp
from routes.api import api_bp
from routes.coding import coding_bp

app.register_blueprint(auth_bp)
app.register_blueprint(home_bp)
app.register_blueprint(admin_bp)
app.register_blueprint(coordinator_bp)
app.register_blueprint(quiz_bp)
app.register_blueprint(study_bp)
app.register_blueprint(api_bp)
app.register_blueprint(coding_bp)

# Session Security & Single-Device Concurrency Middleware
from session_manager import validate_user_session, is_inactive, has_active_exam_session, INACTIVITY_TIMEOUT_SECONDS

@app.before_request
def enforce_session_security():
    path = request.path
    # Allow static files, authentication routes, and public endpoints
    if path.startswith('/static') or path in ['/login', '/logout', '/register', '/api/platform_stats', '/change-password', '/api/auth/change-password', '/api/auth/password-change-required']:
        return None

    user_id = session.get('user_id')

    # Security rule: Force password change on first login before accessing anything else
    if user_id and session.get('must_change_password') == 1:
        if path not in ['/change-password', '/logout', '/api/auth/change-password', '/api/auth/password-change-required']:
            if path.startswith('/api/') or (request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html):
                return jsonify({
                    'success': False,
                    'error': 'password_change_required',
                    'message': 'Password change required before accessing platform features.',
                    'redirect': '/change-password'
                }), 403
            return redirect('/change-password')

    role = session.get('role')

    # Security policy specifically for Student and Coordinator roles
    if role in ['Student', 'Coordinator'] and user_id:
        session_token = session.get('session_token')

        is_api = path.startswith('/api/') or (request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html)

        # 1. Single-Device Concurrency Check:
        # If user logged in on another device, this session token has been superseded.
        if session_token and not validate_user_session(user_id, session_token):
            session.clear()
            if is_api:
                return jsonify({
                    'success': False,
                    'error': 'concurrent_login',
                    'message': 'You have been logged out because your account was logged in on another device.',
                    'redirect': '/login?reason=concurrent_login'
                }), 401
            flash('You have been logged out because your account was logged in on another device.', 'warning')
            return redirect('/login?reason=concurrent_login')

        # 2. Inactivity Timeout Check (25 minutes):
        # Auto logout after 25 minutes if without an active session
        last_activity = session.get('last_activity')
        if last_activity and is_inactive(last_activity):
            # Check if user has an active examination attempt in progress
            if not has_active_exam_session(user_id):
                session.clear()
                if is_api:
                    return jsonify({
                        'success': False,
                        'error': 'session_timeout',
                        'message': 'You have been logged out after 25 minutes of inactivity.',
                        'redirect': '/login?reason=timeout'
                    }), 401
                flash('You have been logged out after 25 minutes of inactivity.', 'info')
                return redirect('/login?reason=timeout')

        # Update last activity timestamp on active request
        session['last_activity'] = time.time()

    return None

@app.after_request
def add_security_headers(response):
    """Applies OWASP standard production security headers to all HTTP responses."""
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    response.headers['X-XSS-Protection'] = '1; mode=block'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    return response

@app.errorhandler(404)
def page_not_found(e):
    if request.path.startswith('/api/') or (request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html):
        return jsonify({'success': False, 'error': 'Endpoint not found', 'status': 404}), 404
    return render_template('error.html', error_title='Page Not Found', error_message='The requested page could not be found.'), 404

@app.errorhandler(500)
def internal_server_error(e):
    if request.path.startswith('/api/') or (request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html):
        return jsonify({'success': False, 'error': 'Internal server error', 'status': 500}), 500
    return render_template('error.html', error_title='Server Error', error_message='An unexpected internal error occurred. Please try again later.'), 500

@app.route('/api/session/heartbeat', methods=['GET', 'POST'])
def session_heartbeat():
    role = session.get('role')
    user_id = session.get('user_id')
    if not user_id or role not in ['Student', 'Coordinator']:
        return jsonify({'active': False}), 401
    session_token = session.get('session_token')
    if session_token and not validate_user_session(user_id, session_token):
        session.clear()
        return jsonify({'active': False, 'error': 'concurrent_login'}), 401
    session['last_activity'] = time.time()
    return jsonify({'active': True, 'timeout_seconds': INACTIVITY_TIMEOUT_SECONDS})

def save_ai_questions():
    from flask import request, redirect, flash
    from utils import collect_selected_ai_preview_rows, get_cached_ai_preview_state
    from database import get_db_connection
    
    preview_token = request.form.get('preview_token')
    quiz_id = request.form.get('quiz_id')
    
    if not quiz_id:
        quiz_id = request.form.get('session_title_select')
    
    preview_state, _ = get_cached_ai_preview_state(preview_token)
    if not preview_state:
        flash('Preview session expired', 'error')
        return redirect('/admin/ai_generate')
    
    selected_rows = collect_selected_ai_preview_rows(request.form)
    
    if not selected_rows:
        flash('No questions selected', 'warning')
        return redirect(f'/admin/ai_preview?preview_token={preview_token}')
    
    if not quiz_id:
        flash('No quiz selected', 'error')
        return redirect('/admin/ai_generate')
    
    conn = get_db_connection()
    added_count = 0
    try:
        with conn.cursor() as cursor:
            for row in selected_rows:
                cursor.execute('''
                    INSERT INTO Questions (quiz_id, question_text, option_a, option_b, option_c, option_d, correct_option, marks)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                ''', (quiz_id, row['question_text'], row['option_a'], row['option_b'], row['option_c'], row['option_d'], row['correct_option'], row['marks']))
                added_count += 1
        conn.commit()
    except Exception as e:
        flash(f'Error saving questions: {str(e)}', 'error')
        return redirect(f'/admin/ai_preview?preview_token={preview_token}')
    finally:
        conn.close()
    
    flash(f'Successfully added {added_count} questions!', 'success')
    return redirect(f'/admin/edit_session/{quiz_id}')

@app.route('/export_ai_preview_docx', methods=['POST'])
def export_ai_preview_docx():
    from flask import request, Response
    from utils import collect_selected_ai_preview_rows, get_cached_ai_preview_state
    from database import get_db_connection
    import io
    from docx import Document
    
    preview_token = request.form.get('preview_token')
    quiz_id = request.form.get('quiz_id') or request.form.get('session_title_select')
    
    preview_state, _ = get_cached_ai_preview_state(preview_token)
    if not preview_state:
        flash('Preview session expired', 'error')
        return redirect('/admin/ai_generate')
    
    selected_rows = collect_selected_ai_preview_rows(request.form)
    
    doc = Document()
    doc.add_heading('AI Generated Questions', 0)
    
    if quiz_id:
        doc.add_paragraph(f'Quiz ID: {quiz_id}')
    if preview_state.get('generation_prompt'):
        doc.add_paragraph(f'Topic: {preview_state["generation_prompt"]}')
    
    for i, row in enumerate(selected_rows, 1):
        doc.add_paragraph(f'{i}. {row["question_text"]}', style='List Number')
        doc.add_paragraph(f'   A) {row["option_a"]}')
        doc.add_paragraph(f'   B) {row["option_b"]}')
        doc.add_paragraph(f'   C) {row["option_c"]}')
        doc.add_paragraph(f'   D) {row["option_d"]}')
        doc.add_paragraph(f'   Answer: {row["correct_option"]} ({row["marks"]} mark(s))')
        doc.add_paragraph('')
    
    f = io.BytesIO()
    doc.save(f)
    f.seek(0)
    
    return Response(
        f.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-disposition': 'attachment; filename=ai_questions.docx'}
    )

# Coordinator dashboard - using blueprint route instead

@app.route('/question/edit/<int:question_id>', methods=['GET', 'POST'])
def edit_question(question_id):
    conn = get_db_connection()
    question = None
    quizzes = []
    
    if conn:
        with conn.cursor() as cursor:
            if request.method == 'POST':
                question_text = request.form.get('question_text')
                option_a = request.form.get('option_a')
                option_b = request.form.get('option_b')
                option_c = request.form.get('option_c')
                option_d = request.form.get('option_d')
                correct_option = request.form.get('correct_option')
                marks = request.form.get('marks', 1)
                module = request.form.get('module', 'General')
                new_quiz_id = request.form.get('quiz_id')
                
                cursor.execute('''
                    UPDATE Questions 
                    SET quiz_id=%s, question_text=%s, option_a=%s, option_b=%s, option_c=%s, 
                        option_d=%s, correct_option=%s, marks=%s, module=%s
                    WHERE question_id=%s
                ''', (new_quiz_id, question_text, option_a, option_b, option_c, option_d, correct_option, marks, module, question_id))
                conn.commit()
                conn.close()
                return redirect(f'/admin/edit_session/{new_quiz_id}')
            
            cursor.execute('SELECT * FROM Questions WHERE question_id=%s', (question_id,))
            question = cursor.fetchone()
            cursor.execute('SELECT quiz_id, title FROM Quizzes ORDER BY title')
            quizzes = cursor.fetchall()
        conn.close()
    
    if not question:
        return "Question not found", 404
    
    return render_template('edit_question.html', q=question, quizzes=quizzes)

@app.route('/session/edit/<int:quiz_id>', methods=['GET', 'POST'])
def alias_edit_session(quiz_id):
    # A 307 Temporary Redirect is used here to ensure that if this is a POST request, 
    # the form payload and the POST method are preserved when forwarding to the correct route.
    return redirect(f'/admin/edit_session/{quiz_id}', code=307)

@app.route('/question/delete/<int:question_id>')
def delete_question(question_id):
    conn = get_db_connection()
    if conn:
        with conn.cursor() as cursor:
            cursor.execute('SELECT quiz_id FROM Questions WHERE question_id=%s', (question_id,))
            result = cursor.fetchone()
            quiz_id = result['quiz_id'] if result else None
            cursor.execute('DELETE FROM Questions WHERE question_id=%s', (question_id,))
            conn.commit()
        conn.close()
        
        if quiz_id:
            return redirect(f'/admin/edit_session/{quiz_id}')
    
    return redirect('/admin/sessions')

@app.route('/upload_docx', methods=['POST'])
def upload_docx():
    from flask import request, jsonify, flash, redirect
    import json
    from routes.admin import parse_questions_from_file

    if 'file' not in request.files:
        if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
            return jsonify({'success': False, 'error': 'No file uploaded'})
        flash('No file uploaded', 'error')
        return redirect('/admin/sessions')

    file = request.files['file']
    if file.filename == '':
        if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
            return jsonify({'success': False, 'error': 'No file selected'})
        flash('No file selected', 'error')
        return redirect('/admin/sessions')

    quiz_id = request.form.get('quiz_id')
    if not quiz_id:
        if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
            return jsonify({'success': False, 'error': 'No assessment session selected'})
        flash('No assessment session selected', 'error')
        return redirect('/admin/sessions')

    fallback_module = request.form.get('module_name', '').strip()

    try:
        valid_q, invalid_q = parse_questions_from_file(file, file.filename)
        if not valid_q:
            err_msg = 'No valid questions could be extracted from this Word document.'
            if invalid_q and invalid_q[0].get('errors'):
                err_msg += ' ' + '; '.join(invalid_q[0]['errors'])
            if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
                return jsonify({'success': False, 'error': err_msg})
            flash(err_msg, 'error')
            return redirect(f'/admin/edit_session/{quiz_id}')

        MODULE_MAP = {
            'mscq_single': 'Single Choice',
            'mscq_multiple': 'Multiple Choice',
            'mscq_select': 'Select Dropdown',
            'fill_blank': 'Fill in the Blanks',
            'true_false': 'True / False',
            'coding': 'Coding'
        }

        conn = get_db_connection()
        questions_added = 0
        with conn.cursor() as cursor:
            for q in valid_q:
                q_text = q.get('question_text_norm') or q.get('questionText') or ''
                q_type = q.get('question_type_norm') or q.get('type') or 'mscq_single'
                marks = q.get('marks_norm', 1)
                neg = q.get('negative_marks_norm', 0.25)
                explanation = q.get('explanation', '')
                correct_ans = q.get('correctAnswer', '')
                options = q.get('options', [])

                mod_name = q.get('module_norm') or q.get('module')
                if not mod_name or mod_name.lower() == 'general':
                    mod_name = fallback_module or MODULE_MAP.get(q_type, 'General')

                meta = {
                    'options': options,
                    'correctAnswer': correct_ans,
                    'negativeMarks': neg,
                    'caseSensitive': bool(q.get('caseSensitive', False)),
                    'acceptedAnswers': q.get('acceptedAnswers') or [correct_ans]
                }

                opt_a = options[0].get('text', '') if len(options) > 0 else ''
                opt_b = options[1].get('text', '') if len(options) > 1 else ''
                opt_c = options[2].get('text', '') if len(options) > 2 else ''
                opt_d = options[3].get('text', '') if len(options) > 3 else ''

                cursor.execute('''
                    INSERT INTO Questions (
                        quiz_id, question_type, question_text,
                        option_a, option_b, option_c, option_d,
                        correct_option, marks, module, subject,
                        difficulty, status, explanation,
                        negative_marks, metadata_json
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ''', (
                    quiz_id, q_type, q_text,
                    opt_a, opt_b, opt_c, opt_d,
                    correct_ans, marks, mod_name, 'General',
                    'Medium', 'Published', explanation,
                    neg, json.dumps(meta)
                ))
                questions_added += 1

            conn.commit()
        conn.close()

        if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
            return jsonify({'success': True, 'count': questions_added, 'message': f'Successfully imported {questions_added} questions.'})
        flash(f'Successfully imported {questions_added} questions across modules.', 'success')
        return redirect(f'/admin/edit_session/{quiz_id}')

    except Exception as e:
        if request.accept_mimetypes.accept_json and not request.accept_mimetypes.accept_html:
            return jsonify({'success': False, 'error': str(e)})
        flash(f'Import error: {str(e)}', 'error')
        return redirect(f'/admin/edit_session/{quiz_id}')

if __name__ == '__main__':
    debug_mode = os.getenv('FLASK_DEBUG', '0') == '1'
    host = os.getenv('FLASK_HOST', '0.0.0.0')
    port = int(os.getenv('PORT', 5000))
    app.run(debug=debug_mode, use_reloader=debug_mode, host=host, port=port)
