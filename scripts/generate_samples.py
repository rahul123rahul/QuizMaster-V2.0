import os
import json
import csv
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_samples():
    os.makedirs('static/samples', exist_ok=True)

    # 1. SAMPLE JSON
    sample_json = [
        {
            "id": 1,
            "type": "mscq",
            "selectionType": "single",
            "questionText": "Which language is used to structure web pages?",
            "options": [
                {"id": "A", "text": "HTML"},
                {"id": "B", "text": "CSS"},
                {"id": "C", "text": "Python"},
                {"id": "D", "text": "SQL"}
            ],
            "correctAnswer": "A",
            "marks": 1,
            "negativeMarks": 0.25,
            "explanation": "HTML is used to structure web pages."
        },
        {
            "id": 2,
            "type": "mscq",
            "selectionType": "multiple",
            "questionText": "Which of the following are programming languages?",
            "options": [
                {"id": "A", "text": "Java"},
                {"id": "B", "text": "Python"},
                {"id": "C", "text": "HTML"},
                {"id": "D", "text": "C++"}
            ],
            "correctAnswer": "A, B, D",
            "marks": 2,
            "negativeMarks": 0.5,
            "explanation": "Java, Python, and C++ are programming languages."
        },
        {
            "id": 3,
            "type": "mscq",
            "selectionType": "select",
            "questionText": "Select the correct database language.",
            "options": [
                {"id": "A", "text": "HTML"},
                {"id": "B", "text": "CSS"},
                {"id": "C", "text": "SQL"},
                {"id": "D", "text": "Python"}
            ],
            "correctAnswer": "C",
            "marks": 1,
            "negativeMarks": 0.25,
            "explanation": "SQL is used to query and manage relational databases."
        },
        {
            "id": 4,
            "type": "fillInTheBlanks",
            "selectionType": "N/A",
            "questionText": "The first letter of the English alphabet is ____.",
            "blanks": [
                {
                    "acceptedAnswers": ["A"]
                }
            ],
            "acceptedAnswers": ["A"],
            "correctAnswer": "A",
            "caseSensitive": True,
            "marks": 1,
            "negativeMarks": 0.25,
            "explanation": "A is the first letter of the English alphabet."
        },
        {
            "id": 5,
            "type": "trueFalse",
            "selectionType": "N/A",
            "questionText": "HTML is a programming language.",
            "options": [
                {"id": "A", "text": "True"},
                {"id": "B", "text": "False"}
            ],
            "correctAnswer": "B",
            "marks": 1,
            "negativeMarks": 0.25,
            "explanation": "HTML is a markup language, not a programming language."
        }
    ]

    with open('static/samples/sample_questions.json', 'w', encoding='utf-8') as f:
        json.dump(sample_json, f, indent=2)
    print("Created static/samples/sample_questions.json")

    # 2. SAMPLE CSV
    with open('static/samples/sample_questions.csv', 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['id', 'type', 'selectionType', 'questionText', 'options', 'correctAnswer', 'marks', 'negativeMarks', 'explanation', 'caseSensitive'])
        writer.writerow(['1', 'mscq', 'single', 'Which language is used to structure web pages?', 'A. HTML | B. CSS | C. Python | D. SQL', 'A', '1', '0.25', 'HTML is used to structure web pages.', 'false'])
        writer.writerow(['2', 'mscq', 'multiple', 'Which of the following are programming languages?', 'A. Java | B. Python | C. HTML | D. C++', 'A, B, D', '2', '0.5', 'Java, Python, and C++ are programming languages.', 'false'])
        writer.writerow(['3', 'mscq', 'select', 'Select the correct database language.', 'A. HTML | B. CSS | C. SQL | D. Python', 'C', '1', '0.25', 'SQL is used to query and manage relational databases.', 'false'])
        writer.writerow(['4', 'fillInTheBlanks', 'N/A', 'The first letter of the English alphabet is ____.', '', 'A', '1', '0.25', 'A is the first letter of the English alphabet.', 'true'])
        writer.writerow(['5', 'trueFalse', 'N/A', 'HTML is a programming language.', 'A. True | B. False', 'B', '1', '0.25', 'HTML is a markup language, not a programming language.', 'false'])
    print("Created static/samples/sample_questions.csv")

    # 3. SAMPLE DOCX
    doc = docx.Document()

    # Title & Introduction
    title_p = doc.add_heading('Question Builder – Word Import Sample', level=0)
    title_p.runs[0].font.color.rgb = RGBColor(79, 70, 229) # Primary indigo color

    intro_p = doc.add_paragraph('Use this Word document as a sample format for importing questions into the Question Builder. The importer should recognize the question type, options, correct answer, marks, negative marks, and explanation.')
    intro_p.runs[0].font.italic = True

    questions_data = [
        {
            "num": 1,
            "rows": [
                ("Type", "mscq"),
                ("Selection Type", "single"),
                ("Question", "Which language is used to structure web pages?"),
                ("Options", "A. HTML\nB. CSS\nC. Python\nD. SQL"),
                ("Correct Answer", "A"),
                ("Marks", "1"),
                ("Negative Marks", "0.25"),
                ("Explanation", "HTML is used to structure web pages.")
            ]
        },
        {
            "num": 2,
            "rows": [
                ("Type", "mscq"),
                ("Selection Type", "multiple"),
                ("Question", "Which of the following are programming languages?"),
                ("Options", "A. Java\nB. Python\nC. HTML\nD. C++"),
                ("Correct Answer", "A, B, D"),
                ("Marks", "2"),
                ("Negative Marks", "0.5"),
                ("Explanation", "Java, Python, and C++ are programming languages.")
            ]
        },
        {
            "num": 3,
            "rows": [
                ("Type", "mscq"),
                ("Selection Type", "select"),
                ("Question", "Select the correct database language."),
                ("Options", "A. HTML\nB. CSS\nC. SQL\nD. Python"),
                ("Correct Answer", "C"),
                ("Marks", "1"),
                ("Negative Marks", "0.25"),
                ("Explanation", "SQL is used to query and manage relational databases.")
            ]
        },
        {
            "num": 4,
            "rows": [
                ("Type", "fillInTheBlanks"),
                ("Selection Type", "N/A"),
                ("Question", "The first letter of the English alphabet is ____."),
                ("Correct Answer", "A"),
                ("Marks", "1"),
                ("Negative Marks", "0.25"),
                ("Explanation", "A is the first letter of the English alphabet."),
                ("Case Sensitive", "true")
            ]
        },
        {
            "num": 5,
            "rows": [
                ("Type", "trueFalse"),
                ("Selection Type", "N/A"),
                ("Question", "HTML is a programming language."),
                ("Options", "A. True\nB. False"),
                ("Correct Answer", "B"),
                ("Marks", "1"),
                ("Negative Marks", "0.25"),
                ("Explanation", "HTML is a markup language, not a programming language.")
            ]
        }
    ]

    for q in questions_data:
        doc.add_heading(f"Question {q['num']}", level=2)
        table = doc.add_table(rows=len(q['rows']), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Light Shading Accent 1' if 'Light Shading Accent 1' in [s.name for s in doc.styles] else 'Table Grid'
        
        for r_idx, (key, val) in enumerate(q['rows']):
            cell_k = table.cell(r_idx, 0)
            cell_v = table.cell(r_idx, 1)
            cell_k.text = key
            cell_v.text = val
            cell_k.paragraphs[0].runs[0].font.bold = True
            cell_k.width = Inches(1.8)
            cell_v.width = Inches(4.5)
        
        doc.add_paragraph() # Spacing

    # Import Rules Section
    rules_h = doc.add_heading('Import Rules', level=2)
    rules = [
        "Supported types: mscq, fillInTheBlanks, trueFalse.",
        "MSCQ selection types: single, multiple, select.",
        "Correct Answer must refer to the option letter for option-based questions.",
        "For multiple-choice questions, separate multiple correct answers with commas.",
        "Marks and Negative Marks must be numeric and greater than or equal to zero.",
        "For Fill in the Blanks, Case Sensitive can be true or false.",
        "The importer must show a preview and validation errors before saving."
    ]
    for rule in rules:
        p = doc.add_paragraph(rule, style='List Bullet')

    doc.save('static/samples/sample_questions.docx')
    print("Created static/samples/sample_questions.docx")

if __name__ == '__main__':
    generate_samples()
